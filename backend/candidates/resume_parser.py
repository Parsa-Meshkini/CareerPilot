"""ResumeParser: deterministic text extraction followed by LLM fact extraction (F02)."""
import io

from agent.llm.client import get_llm_client
from careerpilot.errors import ValidationFailed

from .models import ProfileFact
from .services import FACT_FIELDS, validate_fact

PDF, DOCX, TXT = "application/pdf", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "text/plain"
SUPPORTED = {".pdf": PDF, ".docx": DOCX, ".txt": TXT}
MAX_BYTES = 5 * 1024 * 1024
MAX_TEXT_CHARS = 15000


class NoTextError(Exception):
    """The file has no extractable text (e.g. a scanned PDF)."""


class ResumeParser:
    def __init__(self, llm=None):
        self.llm = llm

    def extract_text(self, data, mime_type):
        if mime_type == PDF:
            import pdfplumber

            with pdfplumber.open(io.BytesIO(data)) as pdf:
                text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        elif mime_type == DOCX:
            import docx

            document = docx.Document(io.BytesIO(data))
            lines = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    lines.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(lines)
        else:
            text = data.decode("utf-8", errors="replace")
        text = "\n".join(line.rstrip() for line in text.splitlines() if line.strip())
        if len(text) < 40:
            raise NoTextError("no readable text found; enter your facts manually (OCR is not supported)")
        return text[:MAX_TEXT_CHARS]

    def extract_facts(self, user, text):
        """Asks the LLM to structure the resume; returns unsaved, unverified ProfileFacts."""
        llm = self.llm or get_llm_client()
        result = llm.generate_structured(user, "resume.extract_facts", {"resume_text": text})
        facts = []
        for item in result.data.get("facts", []):
            fact_type = str(item.get("fact_type", "")).upper()
            if fact_type not in FACT_FIELDS:
                continue
            data = item.get("data") or {}
            fact = ProfileFact(user=user, fact_type=fact_type, source=ProfileFact.RESUME_EXTRACTION, verified=False)
            try:
                fact.data = validate_fact(fact_type, data)
            except ValidationFailed as exc:
                required = FACT_FIELDS[fact_type][0]
                if not any(str(data.get(r) or "").strip() for r in required):
                    continue  # nothing usable
                # keep it, but the student must fix it before it can be verified
                fact.data = {k: v for k, v in data.items() if isinstance(v, (str, list))}
                fact.needs_review = True
                fact.review_note = "; ".join(f"{k}: {v}" for k, v in exc.details.get("fields", {}).items())[:200]
            facts.append(fact)
        return facts
