"""DocumentRenderer: turns structured document content into DOCX, PDF and TXT bytes."""
import io

from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer


class DocumentRenderer:
    """Input is a neutral layout: a list of blocks
    {"kind": "title"|"subtitle"|"heading"|"entry"|"bullets"|"paragraph"|"text", ...}."""

    def render_docx(self, blocks):
        import docx
        from docx.shared import Pt

        d = docx.Document()
        style = d.styles["Normal"]
        style.font.name, style.font.size = "Calibri", Pt(10.5)
        for b in blocks:
            k = b["kind"]
            if k == "title":
                d.add_heading(b["text"], level=0)
            elif k == "subtitle":
                d.add_paragraph(b["text"]).italic = True
            elif k == "heading":
                d.add_heading(b["text"], level=2)
            elif k == "entry":
                p = d.add_paragraph()
                p.add_run(b["text"]).bold = True
                if b.get("right"):
                    p.add_run(f"    {b['right']}")
            elif k == "bullets":
                for item in b["items"]:
                    d.add_paragraph(item, style="List Bullet")
            else:
                d.add_paragraph(b["text"])
        buf = io.BytesIO()
        d.save(buf)
        return buf.getvalue()

    def render_pdf(self, blocks):
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=LETTER, leftMargin=0.7 * inch, rightMargin=0.7 * inch,
                                topMargin=0.6 * inch, bottomMargin=0.6 * inch)
        ss = getSampleStyleSheet()
        body = ParagraphStyle("body", parent=ss["Normal"], fontSize=10, leading=13)
        styles = {"title": ParagraphStyle("t", parent=ss["Title"], fontSize=18, spaceAfter=2),
                  "subtitle": ParagraphStyle("s", parent=body, alignment=1, textColor="#444444"),
                  "heading": ParagraphStyle("h", parent=ss["Heading3"], spaceBefore=8, spaceAfter=2),
                  "entry": ParagraphStyle("e", parent=body, fontName="Helvetica-Bold", spaceBefore=4)}
        story = []
        for b in blocks:
            k = b["kind"]
            if k == "bullets":
                story.append(ListFlowable([ListItem(Paragraph(_esc(i), body), leftIndent=10) for i in b["items"]],
                                          bulletType="bullet", start="•", leftIndent=12))
            elif k == "entry":
                right = f"  <font name='Helvetica' color='#555555'>{_esc(b.get('right', ''))}</font>" if b.get("right") else ""
                story.append(Paragraph(_esc(b["text"]) + right, styles["entry"]))
            else:
                story.append(Paragraph(_esc(b["text"]), styles.get(k, body)))
                if k in ("paragraph", "text"):
                    story.append(Spacer(1, 6))
        doc.build(story)
        return buf.getvalue()

    def render_txt(self, blocks):
        lines = []
        for b in blocks:
            if b["kind"] == "bullets":
                lines.extend(f"  - {i}" for i in b["items"])
            elif b["kind"] == "heading":
                lines.extend(["", b["text"].upper()])
            elif b["kind"] == "entry":
                lines.append(b["text"] + (f" ({b['right']})" if b.get("right") else ""))
            else:
                lines.append(b["text"])
        return ("\n".join(lines).strip() + "\n").encode()


def _esc(text):
    return (text or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
